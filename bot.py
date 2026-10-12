import asyncio
import logging
import os
import re
import io
import json
from pathlib import Path
from datetime import datetime
from urllib.parse import urljoin

import aiohttp
from aiogram import Bot, Dispatcher, F
from aiogram.filters import CommandStart, Command
from aiogram.types import Message, CallbackQuery, BufferedInputFile
from aiogram.utils.keyboard import InlineKeyboardBuilder
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from bs4 import BeautifulSoup

# Load local .env without requiring extra packages. Keep .env private; never upload it.
def load_local_env():
    env_path = Path(".env")
    if not env_path.exists():
        return
    for raw_line in env_path.read_text("utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        os.environ.setdefault(key.strip(), value.strip().strip('"').strip("'"))

load_local_env()
TOKEN = os.getenv("BOT_TOKEN", "").strip()
SITE_URL = "https://politex59.ru/расписание-занятий-ул-юбилейная/"
DATA_DIR = Path("data")
DATA_DIR.mkdir(exist_ok=True)
USERS_FILE = DATA_DIR / "users.json"
LINKS_FILE = Path("schedule_links.json")
ADMIN_IDS = {int(x) for x in os.getenv("ADMIN_IDS", "").split(",") if x.strip().isdigit()}
logging.basicConfig(level=logging.INFO)
log = logging.getLogger("schedule-bot")

# You can add known public Cloud.Mail.ru links in schedule_links.json.
# Keys are weekday names: Понедельник, Вторник, Среда, Четверг, Пятница, Суббота.
DAY_NAMES = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота"]
users = {}
if USERS_FILE.exists():
    try:
        users = json.loads(USERS_FILE.read_text("utf-8"))
    except Exception:
        users = {}

def save_users():
    USERS_FILE.write_text(json.dumps(users, ensure_ascii=False, indent=2), "utf-8")

def clean(v):
    if v is None:
        return ""
    return str(v).replace("\xa0", " ").strip()

def get_font(size, bold=False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()

def find_groups(ws):
    found = {}
    for row in ws.iter_rows():
        for cell in row:
            value = clean(cell.value)
            # Typical college group codes: РМ-23-1, ГЭС-26-2, БС-26.
            if re.fullmatch(r"[А-ЯЁA-Z]{1,5}-\d{2}(?:-\d{1,2})?", value):
                found.setdefault(value, (cell.row, cell.column))
    return found

def read_schedule(path, group):
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        # Prefer sheet that contains the requested group.
        selected = None
        for ws in wb.worksheets:
            groups = find_groups(ws)
            if group in groups:
                selected = (ws, groups[group])
                break
        if not selected:
            return None, "Не нашёл группу в таблице. Проверь написание или загрузи правильный файл."
        ws, (header_row, col) = selected
        lessons = []
        current_course = ""
        # Course labels appear in column A or merged areas around the group header.
        for r in range(1, ws.max_row + 1):
            row_text = " ".join(clean(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 10) + 1))
            course_match = re.search(r"([IVX]+)\s*Курс", row_text, re.I)
            if course_match:
                current_course = course_match.group(1).upper() + " курс"
            # lesson number is in column A, with subject / room / teacher on following rows
            number = clean(ws.cell(r, 1).value)
            if number.isdigit() and 1 <= int(number) <= 10:
                subject = clean(ws.cell(r, col).value)
                room = clean(ws.cell(r + 1, col).value) if r + 1 <= ws.max_row else ""
                teacher = clean(ws.cell(r + 2, col).value) if r + 2 <= ws.max_row else ""
                # Skip blank lesson slots.
                if subject and subject not in ("-", "—"):
                    # The sheet merges each period number over two rows, so the
                    # row carrying period 2/4/6/8/10 can actually be the previous
                    # lesson's teacher row. Do not mistake teacher names for subjects.
                    if re.fullmatch(r"[А-ЯЁA-Z][а-яёa-z-]+\s+(?:[А-ЯЁA-Z]\.){1,2}", subject):
                        continue
                    if subject.lower().startswith("с ") and not room:
                        continue
                    lessons.append({
                        "number": int(number),
                        "subject": subject,
                        "room": room if room and room not in ("-", "—") else "",
                        "teacher": teacher if teacher and teacher not in ("-", "—") else "",
                    })
        date_text = ""
        for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), max_col=min(4, ws.max_column)):
            for cell in row:
                if hasattr(cell.value, "strftime"):
                    date_text = cell.value.strftime("%d.%m.%Y")
        return {"group": group, "date": date_text, "course": current_course, "lessons": lessons}, None
    finally:
        wb.close()

def subject_is_course_label(text):
    return bool(re.search(r"\bКурс\b", text, re.I))

def draw_wrapped(draw, text, font, max_width):
    words = text.split()
    lines, current = [], ""
    for word in words:
        test = (current + " " + word).strip()
        if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [""]

def render_image(schedule, day):
    width = 1000
    margin = 42
    title_font = get_font(42, True)
    subtitle_font = get_font(25)
    subject_font = get_font(25, True)
    detail_font = get_font(20)
    small_font = get_font(18)
    lessons = schedule["lessons"]
    # Image height based on number of populated periods.
    height = 210 + max(1, len(lessons)) * 150 + 45
    im = Image.new("RGB", (width, height), "#F3F6FB")
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((18, 18, width-18, height-18), radius=28, fill="#FFFFFF", outline="#DCE4F0", width=3)
    d.rounded_rectangle((18, 18, width-18, 145), radius=28, fill="#174A8B")
    d.rectangle((18, 105, width-18, 145), fill="#174A8B")
    d.text((margin, 37), "РАСПИСАНИЕ ЗАНЯТИЙ", font=title_font, fill="white")
    date = schedule["date"] or day
    subtitle = f"{schedule['group']}  •  {date}"
    d.text((margin+2, 94), subtitle, font=subtitle_font, fill="#E4EEFF")
    y = 165
    if not lessons:
        d.text((margin, y), "Занятий в таблице не найдено.", font=subject_font, fill="#28364A")
    for item in lessons:
        card_h = 130
        d.rounded_rectangle((margin-5, y, width-margin+5, y+card_h), radius=18, fill="#F7F9FD", outline="#E0E7F1", width=2)
        d.rounded_rectangle((margin+10, y+17, margin+76, y+83), radius=15, fill="#DCEAFE")
        num = str(item["number"])
        bbox = d.textbbox((0, 0), num, font=get_font(28, True))
        d.text((margin+43-(bbox[2]-bbox[0])/2, y+32), num, font=get_font(28, True), fill="#174A8B")
        x = margin + 100
        maxw = width - margin - x - 22
        lines = draw_wrapped(d, item["subject"], subject_font, maxw)
        d.text((x, y+15), lines[0], font=subject_font, fill="#18263A")
        if len(lines) > 1:
            d.text((x, y+45), lines[1], font=subject_font, fill="#18263A")
        details = "  •  ".join(v for v in [item["room"], item["teacher"]] if v)
        d.text((x, y+83), details[:95], font=detail_font, fill="#5A687B")
        y += card_h + 12
    d.text((margin, height-47), "Краевой политехнический колледж", font=small_font, fill="#7C8798")
    bio = io.BytesIO()
    im.save(bio, format="PNG", optimize=True)
    return bio.getvalue()

async def fetch_site_links():
    """Best-effort discovery of public schedule links from the college page."""
    timeout = aiohttp.ClientTimeout(total=25)
    headers = {"User-Agent": "Mozilla/5.0 (compatible; CollegeScheduleBot/1.0)"}
    async with aiohttp.ClientSession(timeout=timeout, headers=headers) as session:
        async with session.get(SITE_URL) as response:
            response.raise_for_status()
            html = await response.text()
    soup = BeautifulSoup(html, "html.parser")
    links = {}
    for a in soup.select("a[href]"):
        label = " ".join(a.get_text(" ", strip=True).split())
        href = urljoin(SITE_URL, a.get("href", ""))
        text = (label + " " + href).lower()
        for day in DAY_NAMES:
            if day.lower() in text and ("cloud.mail.ru" in href or href.lower().endswith((".xlsx", ".xls"))):
                links.setdefault(day, href)
    return links

async def resolve_cloud_download(url):
    """Attempt to resolve a Cloud.Mail.ru public share to a downloadable file."""
    if url.lower().endswith((".xlsx", ".xls")):
        return url
    timeout = aiohttp.ClientTimeout(total=30)
    headers = {"User-Agent": "Mozilla/5.0"}
    async with aiohttp.ClientSession(timeout=timeout, headers=headers) as session:
        async with session.get(url, allow_redirects=True) as response:
            html = await response.text(errors="ignore")
            soup = BeautifulSoup(html, "html.parser")
            # Some public pages expose direct file URLs in metadata.
            for tag in soup.select('meta[content]'):
                content = tag.get("content", "")
                if content.startswith("http") and content.lower().split("?")[0].endswith((".xlsx", ".xls")):
                    return content
            m = None
            if m:
                return m.group(0).replace("\\/", "/").replace("\\u0026", "&")
    return None

async def download_excel(url, destination):
    direct = await resolve_cloud_download(url)
    if not direct:
        raise RuntimeError("Не удалось автоматически получить Excel из Облака Mail.ru. Попроси администратора загрузить файл в бота командой /upload.")
    timeout = aiohttp.ClientTimeout(total=60)
    async with aiohttp.ClientSession(timeout=timeout, headers={"User-Agent": "Mozilla/5.0"}) as session:
        async with session.get(direct, allow_redirects=True) as response:
            response.raise_for_status()
            data = await response.read()
    if not data.startswith(b"PK") and not data.startswith(b"\xd0\xcf"):
        raise RuntimeError("Ссылка не вернула Excel-файл. Администратор может загрузить файл командой /upload.")
    destination.write_bytes(data)
    return destination

def load_links():
    if LINKS_FILE.exists():
        try:
            return json.loads(LINKS_FILE.read_text("utf-8"))
        except Exception:
            pass
    return {}

def main_keyboard():
    kb = InlineKeyboardBuilder()
    kb.button(text="📅 Сегодня", callback_data="day:today")
    kb.button(text="➡️ Завтра", callback_data="day:tomorrow")
    kb.button(text="🗓 Вся неделя", callback_data="week")
    kb.button(text="👥 Выбрать группу", callback_data="choose_group")
    kb.adjust(2, 2)
    return kb.as_markup()

def group_keyboard():
    groups = ["ГЭС-26-1", "ГЭС-26-2", "РМ-26-1", "РМ-26-2", "БС-26", "ТЭ-26", "СГ-26",
              "ГЭС-25-1", "ГЭС-25-2", "РМ-25-1", "РМ-25-2", "БС-25", "ТЭ-25", "СГ-25",
              "СД-25", "РМ-24-1", "РМ-24-2", "БС-24", "ТЭ-24", "СГ-24", "СД-24",
              "РМ-23-1", "РМ-23-2", "БС-23", "ТЭ-23", "СГ-23", "АТ-23"]
    kb = InlineKeyboardBuilder()
    for group in groups:
        kb.button(text=group, callback_data=f"group:{group}")
    kb.adjust(3)
    return kb.as_markup()

def day_for_offset(offset):
    from datetime import date, timedelta
    d = date.today() + timedelta(days=offset)
    # College schedule files may be weekday-specific.
    return DAY_NAMES[d.weekday()] if d.weekday() < len(DAY_NAMES) else None

async def get_file_for_day(day):
    links = load_links()
    link = links.get(day)
    if not link:
        # Try discovering the page's public links.
        try:
            found = await fetch_site_links()
            links.update(found)
            LINKS_FILE.write_text(json.dumps(links, ensure_ascii=False, indent=2), "utf-8")
            link = links.get(day)
        except Exception as e:
            log.warning("Could not scan schedule page: %s", e)
    if link:
        return await download_excel(link, DATA_DIR / f"{day}.xlsx")
    # Fall back to admin-uploaded files named after weekdays.
    local = DATA_DIR / f"{day}.xlsx"
    if local.exists():
        return local
    raise RuntimeError(f"Нет файла расписания на день «{day}». Администратор может загрузить Excel через /upload.")

async def send_schedule(message_or_call, user_id, day):
    group = users.get(str(user_id), {}).get("group")
    if not group:
        text = "Сначала выбери свою группу."
        if isinstance(message_or_call, CallbackQuery):
            await message_or_call.message.answer(text, reply_markup=group_keyboard())
        else:
            await message_or_call.answer(text, reply_markup=group_keyboard())
        return
    try:
        path = await get_file_for_day(day)
        schedule, error = read_schedule(path, group)
        if error:
            raise RuntimeError(error)
        image = render_image(schedule, day)
        caption = f"📅 {day} · {group}" + (f" · {schedule['date']}" if schedule["date"] else "")
        target = message_or_call.message if isinstance(message_or_call, CallbackQuery) else message_or_call
        await target.answer_photo(BufferedInputFile(image, filename=f"{group}_{day}.png"), caption=caption, reply_markup=main_keyboard())
    except Exception as e:
        target = message_or_call.message if isinstance(message_or_call, CallbackQuery) else message_or_call
        await target.answer(f"Не получилось получить расписание: {e}", reply_markup=main_keyboard())

async def main():
    if not TOKEN:
        raise SystemExit("Укажи BOT_TOKEN в переменных окружения или в .env (см. README).")
    bot = Bot(TOKEN)
    dp = Dispatcher()

    @dp.message(CommandStart())
    async def start(message: Message):
        user = users.setdefault(str(message.from_user.id), {})
        save_users()
        if user.get("group"):
            await message.answer(f"Привет! Твоя группа: {user['group']}.", reply_markup=main_keyboard())
        else:
            await message.answer("Привет! Я отправляю расписание колледжа отдельной картинкой для твоей группы. Сначала выбери группу:", reply_markup=group_keyboard())

    @dp.message(Command("help"))
    async def help_cmd(message: Message):
        await message.answer("Команды:\n/start — начать\n/group — сменить группу\n/today — расписание на сегодня\n/tomorrow — на завтра\n/upload — загрузить Excel (только администратор)\n\nТакже используй кнопки меню.", reply_markup=main_keyboard())

    @dp.message(Command("group"))
    async def change_group(message: Message):
        await message.answer("Выбери свою группу:", reply_markup=group_keyboard())

    @dp.callback_query(F.data == "choose_group")
    async def choose_group(call: CallbackQuery):
        await call.message.answer("Выбери свою группу:", reply_markup=group_keyboard())
        await call.answer()

    @dp.callback_query(F.data.startswith("group:"))
    async def set_group(call: CallbackQuery):
        group = call.data.split(":", 1)[1]
        users.setdefault(str(call.from_user.id), {})["group"] = group
        save_users()
        await call.message.answer(f"Готово! Запомнил группу {group}.", reply_markup=main_keyboard())
        await call.answer()

    @dp.callback_query(F.data.startswith("day:"))
    async def choose_day(call: CallbackQuery):
        mode = call.data.split(":", 1)[1]
        offset = 0 if mode == "today" else 1
        day = day_for_offset(offset)
        if not day:
            await call.message.answer("На воскресенье расписание обычно не публикуется.")
        else:
            await send_schedule(call, call.from_user.id, day)
        await call.answer()

    @dp.callback_query(F.data == "week")
    async def week(call: CallbackQuery):
        group = users.get(str(call.from_user.id), {}).get("group")
        if not group:
            await call.message.answer("Сначала выбери группу:", reply_markup=group_keyboard())
        else:
            for day in DAY_NAMES:
                await send_schedule(call, call.from_user.id, day)
        await call.answer()

    @dp.message(Command("today"))
    async def today(message: Message):
        day = day_for_offset(0)
        if day:
            await send_schedule(message, message.from_user.id, day)

    @dp.message(Command("tomorrow"))
    async def tomorrow(message: Message):
        day = day_for_offset(1)
        if day:
            await send_schedule(message, message.from_user.id, day)

    @dp.message(Command("upload"))
    async def upload_prompt(message: Message):
        if ADMIN_IDS and message.from_user.id not in ADMIN_IDS:
            await message.answer("Эта команда доступна только администратору.")
            return
        await message.answer("Отправь Excel-файл документом и подпиши его названием дня: Понедельник, Вторник, Среда, Четверг, Пятница или Суббота.")

    @dp.message(F.document)
    async def receive_file(message: Message):
        if ADMIN_IDS and message.from_user.id not in ADMIN_IDS:
            return
        doc = message.document
        filename = (doc.file_name or "").lower()
        if not filename.endswith((".xlsx", ".xlsm", ".xls")):
            await message.answer("Нужен Excel-файл (.xlsx).")
            return
        caption = (message.caption or "").lower()
        day = next((d for d in DAY_NAMES if d.lower() in caption or d.lower() in filename), None)
        if not day:
            await message.answer("Напиши день в подписи к файлу, например: Вторник.")
            return
        tg_file = await bot.get_file(doc.file_id)
        path = DATA_DIR / f"{day}.xlsx"
        await bot.download_file(tg_file.file_path, destination=path)
        await message.answer(f"Файл расписания на {day} сохранён.")

    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
